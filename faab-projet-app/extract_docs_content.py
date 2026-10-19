import os
import json
import mammoth
from pptx import Presentation

# Paths
KIT_DIR = "/Users/koffiefabienmichael/Downloads/FAAB APPLI/FAAB PROJET/KIT CHEF PROJET"
OUTPUT_JSON = "/Users/koffiefabienmichael/Downloads/FAAB APPLI/FAAB PROJET/faab-projet-app/public/docs_content.json"

TABS_FILES = [
    "Workflow Change Request EN", 
    "Workflow demande de modification FR"
]

content_data = {}

def clean_text_key(text):
    if not text: return ""
    return text.lower().replace("é", "e").replace("è", "e").replace("ê", "e") \
                       .replace("ç", "e").replace("à", "a") \
                       .replace(" ", "").replace("'", "").replace("’", "").replace("-", "") \
                       .replace("_", "").replace(".", "")

def extract_word_mammoth(filepath):
    try:
        with open(filepath, "rb") as docx_file:
            result = mammoth.convert_to_html(docx_file)
            html = result.value
            logo_url = "https://blog-gestion-de-projet.com/wp-content/uploads/2020/07/logo-blog2-293px-transparent.png"
            html = html.replace(f'<img alt="{logo_url}" src="data:image/png;base64,', '<img style="display:none;" src="')
            html = html.replace(logo_url, "")
            return html
    except Exception as e:
        return f"<p style='color:red;'>Mammoth Error: {str(e)}</p>"

def extract_ppt_full(filepath):
    try:
        prs = Presentation(filepath)
        slides = []
        for i, slide in enumerate(prs.slides):
            slide_data = {"id": i+1, "blocks": []}
            if slide.shapes.title:
                slide_data["blocks"].append({"type": "title", "text": slide.shapes.title.text})
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    if shape != slide.shapes.title:
                        slide_data["blocks"].append({"type": "text", "text": shape.text})
                if shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        table_rows.append([cell.text_frame.text.strip() if cell.text_frame else cell.text.strip() for cell in row.cells])
                    slide_data["blocks"].append({"type": "table", "data": table_rows})
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                if notes:
                    slide_data["notes"] = notes
            slides.append(slide_data)
        return slides
    except Exception as e:
        return [{"id": 0, "blocks": [{"type": "text", "text": f"PPT Error: {str(e)}"} ]}]

# Gather all files
all_files = []
for root, _, files in os.walk(KIT_DIR):
    for f in files:
        if not f.startswith('~') and (f.endswith(".docx") or f.endswith(".pptx")):
            all_files.append((f, os.path.join(root, f)))

# Extract only target files
for target_name in TABS_FILES:
    target_cleaned = clean_text_key(target_name)
    for fname, fpath in all_files:
        if target_cleaned in clean_text_key(fname):
            key = target_name.lower()
            if fpath.endswith(".docx"):
                content_data[key] = { "type": "word", "content": extract_word_mammoth(fpath) }
            else:
                content_data[key] = { "type": "ppt", "slides": extract_ppt_full(fpath) }
            break

with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
    json.dump(content_data, f, ensure_ascii=False, indent=2)

print(f"Extraction successful: {len(content_data)} files.")
